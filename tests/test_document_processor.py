"""
test_document_processor.py – Tests for document_processor.py

Covers Phase 2 requirements:
  • Validation (extension, magic bytes, size limit, empty files)
  • PDF extraction via PyMuPDF with page metadata
  • DOCX extraction via python-docx with paragraphs and tables
  • Scanned PDF detection reporting OCR requirement
  • Whitespace normalization and extraction artifact cleaning
  • Corrupt and malformed file handling
  • Typed data structures (DocumentMetadata, PageContent, ExtractedDocument)
"""

import io
import pytest
import pymupdf as fitz
from docx import Document

from document_processor import (
    CorruptDocumentError,
    DocumentMetadata,
    DocumentProcessingError,
    EmptyDocumentError,
    ExtractedDocument,
    PageContent,
    ScannedDocumentError,
    UnsupportedDocumentError,
    clean_extraction_artifacts,
    extract_docx_document,
    extract_pdf_document,
    extract_text,
    process_document,
    validate_file,
)


class TestValidateFile:
    def test_valid_pdf(self):
        assert validate_file("contract.pdf", 1024, b"%PDF-1.4 test") is None

    def test_valid_docx(self):
        assert validate_file("contract.docx", 1024, b"PK\x03\x04 test") is None

    def test_unsupported_extension(self):
        err = validate_file("contract.txt", 1024)
        assert err is not None
        assert "Unsupported" in err

    def test_file_too_large(self):
        err = validate_file("contract.pdf", 100 * 1024 * 1024)
        assert err is not None
        assert "limit" in err.lower()

    def test_case_insensitive_extension(self):
        assert validate_file("CONTRACT.PDF", 1024, b"%PDF-1.4") is None
        assert validate_file("Contract.DOCX", 1024, b"PK\x03\x04") is None

    def test_empty_file_validation(self):
        err = validate_file("contract.pdf", 0)
        assert err is not None
        assert "empty" in err.lower()

    def test_invalid_magic_bytes_pdf(self):
        err = validate_file("fake.pdf", 1024, b"NOT_A_PDF_HEADER")
        assert err is not None
        assert "Invalid PDF header" in err

    def test_invalid_magic_bytes_docx(self):
        err = validate_file("fake.docx", 1024, b"NOT_A_DOCX_HEADER")
        assert err is not None
        assert "Invalid DOCX header" in err


class TestExtractPdf:
    def test_single_page_pdf(self):
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "Master Services Agreement\n1. Payment Terms\nNet 30 days.")
        pdf_bytes = doc.tobytes()
        doc.close()

        extracted = process_document("agreement.pdf", pdf_bytes)
        assert isinstance(extracted, ExtractedDocument)
        assert extracted.metadata.file_type == "pdf"
        assert extracted.metadata.page_count == 1
        assert "Master Services Agreement" in extracted.normalized_text
        assert len(extracted.pages) == 1
        assert extracted.pages[0].page_number == 1

    def test_multi_page_pdf(self):
        doc = fitz.open()
        p1 = doc.new_page()
        p1.insert_text((50, 50), "Page 1: 1. Definitions\nConfidential Information definition.")
        p2 = doc.new_page()
        p2.insert_text((50, 50), "Page 2: 2. Termination\nEither party may terminate on 30 days notice.")
        pdf_bytes = doc.tobytes()
        doc.close()

        extracted = process_document("multi_page.pdf", pdf_bytes)
        assert extracted.metadata.page_count == 2
        assert len(extracted.pages) == 2
        assert "Definitions" in extracted.pages[0].text
        assert "Termination" in extracted.pages[1].text

    def test_empty_pdf_document(self):
        doc = fitz.open()
        doc.new_page()  # Page without text
        pdf_bytes = doc.tobytes()
        doc.close()

        with pytest.raises(EmptyDocumentError):
            extract_pdf_document("empty.pdf", pdf_bytes)

    def test_corrupt_pdf(self):
        corrupt_bytes = b"%PDF-1.4 but then total gibberish corrupt content that cannot parse"
        with pytest.raises(CorruptDocumentError):
            extract_pdf_document("corrupted.pdf", corrupt_bytes)


class TestScannedPdfDetection:
    def test_scanned_pdf_reports_ocr_needed(self):
        doc = fitz.open()
        page = doc.new_page()
        # Insert image without any selectable text
        pix = fitz.Pixmap(fitz.csRGB, (0, 0, 10, 10), 1)
        page.insert_image(fitz.Rect(0, 0, 100, 100), stream=pix.tobytes("png"))
        pdf_bytes = doc.tobytes()
        doc.close()

        with pytest.raises(ScannedDocumentError) as exc_info:
            extract_pdf_document("scanned.pdf", pdf_bytes)
        assert "OCR is required" in str(exc_info.value)


class TestExtractDocx:
    def test_docx_paragraphs(self):
        doc = Document()
        doc.add_paragraph("Master Consulting Agreement")
        doc.add_paragraph("Section 1. Scope of Services.")
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        extracted = process_document("consulting.docx", docx_bytes)
        assert isinstance(extracted, ExtractedDocument)
        assert extracted.metadata.file_type == "docx"
        assert "Master Consulting Agreement" in extracted.normalized_text
        assert "Scope of Services" in extracted.normalized_text

    def test_docx_tables(self):
        doc = Document()
        doc.add_paragraph("Fee Schedule")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Milestone 1"
        table.cell(0, 1).text = "$10,000"
        table.cell(1, 0).text = "Milestone 2"
        table.cell(1, 1).text = "$15,000"
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        extracted = process_document("table_test.docx", docx_bytes)
        assert "Milestone 1" in extracted.normalized_text
        assert "$10,000" in extracted.normalized_text

    def test_empty_docx_raises_empty_error(self):
        doc = Document()
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        with pytest.raises(EmptyDocumentError):
            extract_docx_document("empty.docx", docx_bytes)

    def test_corrupt_docx_raises_corrupt_error(self):
        # A docx with PK header but corrupted zip archive
        bad_docx = b"PK\x03\x04corrupted_zip_payload_cannot_extract"
        with pytest.raises(CorruptDocumentError):
            extract_docx_document("corrupted.docx", bad_docx)


class TestArtifactCleaningAndNormalization:
    def test_hyphenated_line_break_rejoining(self):
        text = "This is an indemni-\n   fication clause with non-compete terms."
        cleaned = clean_extraction_artifacts(text)
        assert "indemnification" in cleaned
        # Intra-line compound words must be preserved
        assert "non-compete" in cleaned

    def test_standalone_page_number_removal(self):
        text = (
            "Clause 1. Definitions\n"
            "Page 1 of 5\n"
            "These are the terms.\n"
            "- 2 -\n"
            "Clause 2. Payment\n"
            "— 3 —\n"
            "Net 30 days."
        )
        cleaned = clean_extraction_artifacts(text)
        assert "Page 1 of 5" not in cleaned
        assert "- 2 -" not in cleaned
        assert "— 3 —" not in cleaned
        assert "Clause 1. Definitions" in cleaned
        assert "Clause 2. Payment" in cleaned

    def test_control_character_removal(self):
        text = "Section 1\x0c\nBody text\x00 with soft\xadhyphen."
        cleaned = clean_extraction_artifacts(text)
        assert "\x0c" not in cleaned
        assert "\x00" not in cleaned
        assert "\xad" not in cleaned
        assert "Body text with softhyphen." in cleaned

    def test_horizontal_whitespace_collapsing(self):
        text = "Party A    agrees\t\tto   pay   fees."
        cleaned = clean_extraction_artifacts(text)
        assert cleaned == "Party A agrees to pay fees."


class TestProcessDocumentDispatcher:
    def test_unsupported_extension_raises_error(self):
        with pytest.raises(UnsupportedDocumentError):
            process_document("test.txt", b"plain text")

    def test_backward_compatible_extract_text(self):
        doc = Document()
        doc.add_paragraph("Backward compatibility test.")
        buf = io.BytesIO()
        doc.save(buf)
        text = extract_text("test.docx", buf.getvalue())
        assert "Backward compatibility test." in text
