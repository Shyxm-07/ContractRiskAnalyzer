"""
test_e2e_pipeline.py – Phase 6 End-to-End Integration, Resilience & Security Tests.

Covers:
1. Full end-to-end pipeline with synthetic PDF contract (PyMuPDF -> split -> classify -> rules -> score -> report).
2. Full end-to-end pipeline with synthetic DOCX contract (python-docx -> split -> classify -> rules -> score -> report).
3. Robust handling of edge cases: empty files, corrupt files, unsupported extensions, size limits.
4. Resilience against model errors and LLM fallback.
5. Security & Privacy verification: in-memory processing, no confidential text in logs.
6. Verification that external LLM calls are off by default and template provider works offline.
"""

import io
import logging
import pytest
from unittest.mock import patch

import fitz  # PyMuPDF
from docx import Document

from classifier import classify_clauses
from clause_splitter import segment_document, split_into_clauses
from config import settings
from document_processor import (
    CorruptDocumentError,
    DocumentProcessingError,
    EmptyDocumentError,
    UnsupportedDocumentError,
    process_document,
    validate_file,
)
from llm_explainer import (
    Explanation,
    explain_flags,
    MockLLMProvider,
    HuggingFaceInferenceProvider,
    TemplateFallbackProvider,
)
from report_generator import generate_pdf_report
from risk_rules import evaluate_clause, evaluate_document, RiskFlag
from risk_scorer import score_contract, OverallRiskReport


# ── Fixtures & Helper Generators ──────────────────────────────────────

def _create_synthetic_pdf_bytes() -> bytes:
    """Create a multi-page synthetic commercial agreement in-memory using PyMuPDF."""
    doc = fitz.open()

    # Page 1: Payment & Termination
    p1 = doc.new_page()
    text_p1 = (
        "MASTER SERVICES AGREEMENT\n\n"
        "1. PAYMENT TERMS\n"
        "Client shall remit payment to Contractor promptly in due course following receipt "
        "of each undisputed invoice.\n\n"
        "2. TERMINATION\n"
        "Company reserves the right to terminate this Agreement immediately at any time "
        "for convenience without notice, liability, or penalty.\n"
    )
    p1.insert_text((50, 72), text_p1, fontsize=11)

    # Page 2: Liability & Indemnity & Governing Law
    p2 = doc.new_page()
    text_p2 = (
        "3. INDEMNIFICATION\n"
        "Vendor shall indemnify, defend, and hold harmless Client against any and all claims, "
        "losses, and damages arising out of this Agreement.\n\n"
        "4. LIMITATION OF LIABILITY\n"
        "Neither party's liability under this Agreement shall be subject to any financial cap.\n\n"
        "5. GOVERNING LAW\n"
        "This Agreement shall be governed by the laws of the State of New York.\n"
    )
    p2.insert_text((50, 72), text_p2, fontsize=11)

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def _create_synthetic_docx_bytes() -> bytes:
    """Create a synthetic commercial agreement in-memory using python-docx."""
    doc = Document()
    doc.add_heading("Software Consulting Services Agreement", level=1)

    doc.add_heading("1. Scope and Deadlines", level=2)
    doc.add_paragraph(
        "Contractor will deliver milestones within a reasonable time as determined by Contractor."
    )

    doc.add_heading("2. Automatic Renewal", level=2)
    doc.add_paragraph(
        "This Agreement will automatically renew indefinitely for successive 1-year terms "
        "unless notice is given 365 days prior."
    )

    doc.add_heading("3. Confidentiality", level=2)
    doc.add_paragraph(
        "Both parties shall maintain strict confidentiality for a period of five (5) years."
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── 1. Full E2E Pipeline Tests ────────────────────────────────────────

class TestEndToEndPipeline:
    def test_e2e_synthetic_pdf_pipeline(self):
        """Full end-to-end execution on a synthetic PDF document."""
        pdf_bytes = _create_synthetic_pdf_bytes()

        # Step 1: Ingestion & Extraction
        extracted_doc = process_document("Commercial_Contract.pdf", pdf_bytes)
        assert extracted_doc.metadata.file_type == "pdf"
        assert extracted_doc.metadata.page_count == 2
        assert extracted_doc.metadata.word_count > 30

        # Step 2: Clause Segmentation
        clauses = segment_document(extracted_doc)
        assert len(clauses) >= 4
        # Verify page number preservation
        assert any(c.page_number == 1 for c in clauses)
        assert any(c.page_number == 2 for c in clauses)

        # Step 3: ML Classification (deterministic baseline)
        classifications = classify_clauses(clauses, backend="tfidf")
        assert len(classifications) == len(clauses)
        assert all(c.predicted_category in settings.clause_categories for c in classifications)

        # Step 4: Legal Risk Detection (clause & doc level)
        clause_flags = {}
        for c in clauses:
            flags = evaluate_clause(c.text)
            if flags:
                clause_flags[c.index] = flags
        assert len(clause_flags) >= 2  # At least payment/timeline & unilateral termination

        doc_flags = evaluate_document(extracted_doc.normalized_text)
        assert isinstance(doc_flags, list)

        # Step 5: Transparent Weighted Scoring
        report: OverallRiskReport = score_contract(
            classifications,
            clause_flags,
            doc_flags,
            profile_name="standard_commercial",
        )
        assert 0.0 <= report.overall_score <= 100.0
        assert report.risk_level in ("low", "medium", "high")
        assert len(report.clause_details) == len(clauses)
        assert report.is_deterministic is True

        # Step 6: GenAI Explanations (Mock LLM offline)
        explanations = []
        for detail in report.clause_details:
            clause_obj = clauses[detail.clause_index]
            expl = explain_flags(
                clause_index=detail.clause_index,
                heading=detail.heading,
                flags=detail.flags,
                clause_text=clause_obj.text,
                use_llm=True,
                provider_type="mock",
                clause_id=clause_obj.clause_id,
            )
            explanations.append(expl)
        assert len(explanations) == len(clauses)

        # Verify structured outputs and grounding on mock LLM explanations
        for expl in explanations:
            if expl.status == "completed_llm" and expl.structured_output:
                assert expl.structured_output.grounding_verified is True
                assert expl.structured_output.suggested_revision != ""
                assert "does not constitute legal advice" in expl.structured_output.disclaimer.lower()

        # Step 7: Executive PDF Report Generation
        pdf_report_bytes = generate_pdf_report(
            report,
            explanations,
            clauses=clauses,
            doc_metadata=extracted_doc.metadata,
        )
        assert isinstance(pdf_report_bytes, bytes)
        assert len(pdf_report_bytes) > 2500
        assert pdf_report_bytes.startswith(b"%PDF")

    def test_e2e_synthetic_docx_pipeline(self):
        """Full end-to-end execution on a synthetic DOCX document."""
        docx_bytes = _create_synthetic_docx_bytes()

        # Step 1: Ingestion
        extracted_doc = process_document("Consulting_Agreement.docx", docx_bytes)
        assert extracted_doc.metadata.file_type == "docx"
        assert extracted_doc.metadata.word_count > 20

        # Step 2: Segmentation
        clauses = segment_document(extracted_doc)
        assert len(clauses) >= 3

        # Step 3: Classification & Risk Detection
        classifications = classify_clauses(clauses, backend="tfidf")
        clause_flags = {
            c.index: evaluate_clause(c.text)
            for c in clauses
            if evaluate_clause(c.text)
        }
        doc_flags = evaluate_document(extracted_doc.normalized_text)

        # Step 4: Scoring
        report = score_contract(
            classifications, clause_flags, doc_flags, profile_name="buyer_protective"
        )
        assert 0.0 <= report.overall_score <= 100.0

        # Step 5: Report Generation with template fallback explainer
        explanations = [
            explain_flags(
                d.clause_index,
                d.heading,
                d.flags,
                clause_text=clauses[d.clause_index].text,
                use_llm=False,
            )
            for d in report.clause_details
        ]
        pdf_bytes = generate_pdf_report(
            report, explanations, clauses=clauses, doc_metadata=extracted_doc.metadata
        )
        assert pdf_bytes.startswith(b"%PDF")

    def test_e2e_deterministic_score_preservation(self):
        """Run pipeline twice on identical input and verify scores match exactly."""
        pdf_bytes = _create_synthetic_pdf_bytes()

        scores = []
        for _ in range(2):
            doc = process_document("Test.pdf", pdf_bytes)
            clauses = segment_document(doc)
            cls = classify_clauses(clauses, backend="tfidf")
            cflags = {c.index: evaluate_clause(c.text) for c in clauses if evaluate_clause(c.text)}
            dflags = evaluate_document(doc.normalized_text)
            report = score_contract(cls, cflags, dflags, profile_name="standard_commercial")
            scores.append(report.overall_score)

        assert scores[0] == scores[1], f"Scores diverged: {scores[0]} vs {scores[1]}"


# ── 2. Error Handling & Edge Cases ────────────────────────────────────

class TestEdgeCasesAndResilience:
    def test_empty_file_raises_error(self):
        """Empty documents must raise EmptyDocumentError."""
        with pytest.raises(EmptyDocumentError):
            process_document("empty.pdf", b"")

    def test_empty_file_validation_returns_error(self):
        """validate_file returns an error string for empty files."""
        result = validate_file("empty.docx", 0, b"")
        assert result is not None
        assert "empty" in result.lower()

    def test_corrupt_pdf_raises_error(self):
        """Corrupt PDF must raise CorruptDocumentError."""
        with pytest.raises(CorruptDocumentError):
            process_document("corrupted.pdf", b"%PDF-1.4 corrupt content that cannot parse")

    def test_corrupt_docx_raises_error(self):
        """Corrupt DOCX must raise CorruptDocumentError."""
        with pytest.raises(CorruptDocumentError):
            process_document("corrupted.docx", b"PK\x03\x04 corrupt zip archive")

    def test_unsupported_file_extension_validation(self):
        """validate_file returns error for unsupported extensions."""
        result = validate_file("contract.exe", 100, b"binarycontent")
        assert result is not None
        assert "unsupported" in result.lower() or "allowed" in result.lower()

    def test_unsupported_extension_process_raises(self):
        """process_document raises UnsupportedDocumentError for unknown types."""
        with pytest.raises(UnsupportedDocumentError):
            process_document("contract.txt", b"plain text")

    def test_file_size_validation(self):
        """validate_file returns error for oversized files."""
        huge_size = 1024 * 1024 * 100  # 100 MB
        result = validate_file("large.pdf", huge_size, b"%PDF-1.4")
        # Should report size issue
        assert result is not None

    def test_hf_provider_fallback_without_api_key(self):
        """HF provider returns fallback status when no API key is configured."""
        provider = HuggingFaceInferenceProvider(api_key="", endpoint_url="")
        output, status, err = provider.explain_clause(
            clause_id="clause_001",
            clause_text="Client shall pay promptly.",
            flags=[
                RiskFlag(
                    rule_id="RULE-DEADLINE-001",
                    label="Indefinite Performance Timeline",
                    severity="medium",
                    description="Timeline lacks fixed days.",
                    matched_snippet="promptly",
                )
            ],
        )
        assert status == "fallback_template"
        assert output is None

    def test_mock_provider_works_offline(self):
        """Mock provider always succeeds with validated structured outputs."""
        provider = MockLLMProvider()
        output, status, err = provider.explain_clause(
            clause_id="clause_test",
            clause_text="Client shall remit payment promptly in due course.",
            flags=[
                RiskFlag(
                    rule_id="RULE-DEADLINE-001",
                    label="Indefinite Performance Timeline",
                    severity="medium",
                    description="Payment deadline is ambiguous.",
                    matched_snippet="promptly in due course",
                )
            ],
        )
        assert status == "completed_llm"
        assert output is not None
        assert output.grounding_verified is True
        assert output.suggested_revision != ""

    def test_template_fallback_provider_no_flags(self):
        """Template fallback produces output even when no flags are present."""
        provider = TemplateFallbackProvider()
        output, status, err = provider.explain_clause(
            clause_id="clause_clean",
            clause_text="Both parties agree to mutual confidentiality obligations.",
            flags=[],
        )
        assert status == "completed_template"
        assert output is not None
        assert "no specific risks" in output.explanation.lower()


# ── 3. Security, Privacy & In-Memory Guarantees ───────────────────────

class TestSecurityAndPrivacy:
    def test_no_confidential_text_in_logs(self, caplog):
        """Log messages must never record raw contract clause text."""
        caplog.set_level(logging.DEBUG)

        secret_phrase = "TOP_SECRET_ACQUISITION_PRICE_USD_999_MILLION"
        clause_text = f"1. Scope of Work\nThe parties agree that {secret_phrase} shall apply."

        # Run pipeline steps
        clauses = split_into_clauses(clause_text)
        classifications = classify_clauses(clauses, backend="tfidf")
        flags = {c.index: evaluate_clause(c.text) for c in clauses}
        doc_flags = evaluate_document(clause_text)
        report = score_contract(classifications, flags, doc_flags)
        expl = explain_flags(0, "Scope", [], clause_text=clause_text, use_llm=False)

        # Inspect all log records
        for record in caplog.records:
            assert secret_phrase not in record.getMessage(), (
                f"Confidential contract text was leaked into log record: {record.getMessage()}"
            )

    def test_external_llm_disabled_by_default(self):
        """External LLM must be disabled by default (enable_external_llm=False)."""
        assert settings.enable_external_llm is False

    def test_default_provider_is_template(self):
        """Default provider type should be template (offline)."""
        assert settings.llm_provider in ("template", "mock")

    def test_explain_flags_offline_no_external_calls(self):
        """Calling explain_flags with use_llm=False must not invoke external APIs."""
        expl = explain_flags(
            0, "Recitals", [], clause_text="Some text", use_llm=False
        )
        # Template-based status
        assert expl.status in ("completed_template", "completed_offline")
        # No LLM structured output (template produces one but it's a template variant)
        assert expl.provider_used == "TemplateFallbackProvider"
